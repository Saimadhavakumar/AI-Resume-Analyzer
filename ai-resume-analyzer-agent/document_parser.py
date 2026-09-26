# document_parser.py
# PDF and DOCX text extraction layer

import os


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file using PyMuPDF (fitz).

    Handles:
    - Multi-page PDFs
    - Empty documents
    - Corrupted files
    - Scanned/image-only PDFs (returns clear message, no OCR)

    Args:
        file_path: Absolute or relative path to the PDF file.

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: If the file is empty, corrupted, or contains no extractable text.
        FileNotFoundError: If the file does not exist.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
    except ImportError:
        raise ImportError(
            "PyMuPDF is required for PDF parsing. Install it with: pip install PyMuPDF"
        )

    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise ValueError(f"Could not open PDF file. The file may be corrupted: {e}")

    if doc.page_count == 0:
        doc.close()
        raise ValueError("The PDF file contains no pages.")

    pages_text = []
    for page_num in range(doc.page_count):
        page = doc[page_num]
        text = page.get_text("text")
        if text and text.strip():
            pages_text.append(text)

    doc.close()

    if not pages_text:
        raise ValueError(
            "No extractable text found in this PDF. "
            "The document may contain only scanned images. "
            "Please provide a PDF with selectable text, or use a DOCX file instead."
        )

    return "\n\n".join(pages_text)


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file using python-docx.

    Handles:
    - Multi-paragraph documents
    - Tables (extracts cell text row by row)
    - Empty documents
    - Corrupted files

    Args:
        file_path: Absolute or relative path to the DOCX file.

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: If the file is empty or corrupted.
        FileNotFoundError: If the file does not exist.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX parsing. Install it with: pip install python-docx"
        )

    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Could not open DOCX file. The file may be corrupted: {e}")

    parts = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)

    if not parts:
        raise ValueError(
            "The DOCX file contains no extractable text. "
            "Please upload a non-empty resume document."
        )

    return "\n".join(parts)


def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    """
    Extract text from file bytes (used by the web upload handler).

    Determines the file type from the filename extension and delegates
    to the appropriate parser. Uses a temporary file for processing.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename: Original filename (used to determine file type).

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: If the file type is unsupported or extraction fails.
    """
    import tempfile

    ext = os.path.splitext(filename)[1].lower()

    if ext not in (".pdf", ".docx"):
        raise ValueError(
            f"Unsupported file type: '{ext}'. "
            "Please upload a PDF (.pdf) or Word (.docx) file."
        )

    # Write bytes to a temporary file for parsing
    suffix = ext
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        if ext == ".pdf":
            return extract_text_from_pdf(tmp_path)
        else:
            return extract_text_from_docx(tmp_path)
    finally:
        # Clean up temp file
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def extract_text(file_path: str) -> str:
    """
    Unified text extraction interface.

    Determines the file type from the extension and delegates
    to the appropriate parser.

    Args:
        file_path: Path to the resume file (PDF or DOCX).

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: If the file type is unsupported or extraction fails.
        FileNotFoundError: If the file does not exist.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(
            f"Unsupported file type: '{ext}'. "
            "Supported formats: PDF (.pdf), Word (.docx)"
        )
